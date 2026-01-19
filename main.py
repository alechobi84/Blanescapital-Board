"""
Blanes Capital Dashboard - Flask Application
Para ejecutar en PythonAnywhere
"""

from flask import Flask, render_template, request, redirect, url_for, flash, session, send_file
from flask_sqlalchemy import SQLAlchemy
from flask_login import LoginManager, UserMixin, login_user, logout_user, login_required, current_user
from werkzeug.security import generate_password_hash, check_password_hash
import os
from io import BytesIO
from datetime import datetime

# ============================================
# CONFIGURACIÓN
# ============================================
app = Flask(__name__)
app.config['SECRET_KEY'] = os.environ.get('SECRET_KEY', 'blanes-capital-secret-key-cambiar-en-produccion')
app.config['SQLALCHEMY_DATABASE_URI'] = os.environ.get('DATABASE_URL', 'sqlite:///blanes.db')
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False

db = SQLAlchemy(app)
login_manager = LoginManager(app)
login_manager.login_view = 'login'
login_manager.login_message = 'Por favor, inicia sesión para acceder.'

# ============================================
# MODELOS DE BASE DE DATOS
# ============================================
class User(UserMixin, db.Model):
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(80), unique=True, nullable=False)
    email = db.Column(db.String(120), unique=True, nullable=True)
    password_hash = db.Column(db.String(256), nullable=False)
    is_admin = db.Column(db.Boolean, default=False)

    def set_password(self, password):
        self.password_hash = generate_password_hash(password)

    def check_password(self, password):
        return check_password_hash(self.password_hash, password)

@login_manager.user_loader
def load_user(user_id):
    return User.query.get(int(user_id))

# ============================================
# RUTAS
# ============================================
@app.route('/')
def index():
    """Página principal - redirige al login si no está autenticado"""
    if current_user.is_authenticated:
        return render_template('index.html')
    return redirect(url_for('login'))

@app.route('/login', methods=['GET', 'POST'])
def login():
    """Página de login"""
    if current_user.is_authenticated:
        return redirect(url_for('dashboard'))

    if request.method == 'POST':
        username = request.form.get('username')
        password = request.form.get('password')
        user = User.query.filter_by(username=username).first()

        if user and user.check_password(password):
            login_user(user)
            next_page = request.args.get('next')
            return redirect(next_page or url_for('dashboard'))
        else:
            flash('Usuario o contraseña incorrectos', 'error')

    return render_template('login.html')

@app.route('/dashboard')
@login_required
def dashboard():
    """Dashboard principal - requiere autenticación"""
    return render_template('index.html')

@app.route('/logout')
@login_required
def logout():
    """Cerrar sesión"""
    logout_user()
    flash('Has cerrado sesión correctamente', 'success')
    return redirect(url_for('login'))

# ============================================
# DATOS PROTEGIDOS (solo accesibles con login)
# ============================================
PERIOD_DATA = {
    '2024': {
        'title': 'Cierre 2024',
        'date': '31 de diciembre de 2024',
        'total': {
            'bajoGestion': 161.06,
            'rentabilidad': 6.36,
            'rentasInmob': 364,
            'costeEstructura': 0.35,
            'costeEstructuraTotal': 547
        },
        'pablo': {
            'patrimonioTotal': 126.09,
            'bajoGestion': 88.43,
            'empresarial': 37.66,
            'rentabilidad': 5.69,
            'rentasInmob': 186,
            'carteraFinanciera': 46.43,
            'carteraInmobiliaria': 28.0,
            'alternativas': 1.68,
            'rotacion': 3.44,
            'exposicionUSD': 30,
            'rf': 78, 'rv': 18, 'mp': 4,
            'bancos': { 'jp': 58, 'gs': 22, 'march': 13, 'andbank': 7 }
        },
        'ale': {
            'patrimonioTotal': 150.08,
            'bajoGestion': 72.63,
            'empresarial': 77.45,
            'rentabilidad': 7.22,
            'rentasInmob': 178,
            'carteraFinanciera': 34.50,
            'carteraInmobiliaria': 21.16,
            'alternativas': 4.37,
            'rotacion': 3.44,
            'exposicionUSD': 32,
            'rf': 52, 'rv': 45, 'mp': 3,
            'bancos': { 'jp': 64, 'gs': 16, 'andbank': 12, 'march': 8 }
        },
        'estructura': {
            'total': 546915,
            'ratio': 0.35,
            'budget': 543740,
            'mensual': {
                'ene': 45, 'feb': 46, 'mar': 42, 'abr': 45, 'may': 44, 'jun': 48,
                'jul': 46, 'ago': 44, 'sep': 45, 'oct': 47, 'nov': 48, 'dic': 47
            }
        }
    },
    '1s2025': {
        'title': '1er Semestre 2025',
        'date': '30 de junio de 2025',
        'total': {
            'bajoGestion': 175.1,
            'rentabilidad': -0.03,
            'rentasInmob': 748,
            'costeEstructura': 0.33,
            'costeEstructuraTotal': 278
        },
        'pablo': {
            'patrimonioTotal': 128.7,
            'bajoGestion': 96.4,
            'empresarial': 32.3,
            'rentabilidad': -0.14,
            'rentabilidadSinUSD': 3.07,
            'rentasInmob': 385,
            'carteraFinanciera': 55.8,
            'carteraInmobiliaria': 28.0,
            'alternativas': 2.2,
            'rotacion': 3.44,
            'exposicionUSD': 29,
            'ocupacion': 79,
            'rf': 75, 'rv': 21, 'mp': 4,
            'bancos': { 'jp': 60, 'gs': 21, 'march': 12, 'andbank': 7 }
        },
        'ale': {
            'patrimonioTotal': 143.4,
            'bajoGestion': 78.8,
            'empresarial': 64.7,
            'rentabilidad': 0.08,
            'rentabilidadSinUSD': 3.86,
            'rentasInmob': 363,
            'carteraFinanciera': 43.0,
            'carteraInmobiliaria': 21.6,
            'alternativas': 5.4,
            'rotacion': 3.44,
            'exposicionUSD': 27,
            'ocupacion': 88,
            'rf': 50, 'rv': 47, 'mp': 2,
            'bancos': { 'jp': 66, 'gs': 14, 'andbank': 12, 'march': 8 }
        },
        'estructura': {
            'total': 278000,
            'ratio': 0.33,
            'budget': 315000,
            'mensual': {
                'ene': 44, 'feb': 48, 'mar': 40, 'abr': 44, 'may': 40, 'jun': 63
            }
        }
    }
}

HISTORICAL_DATA = {
    'years': ['2020', '2021', '2022', '2023', '2024', '1S 2025'],
    'pablo': {
        'bajoGestion': [111.9, 86.5, 85.0, 87.9, 88.43, 96.4],
        'empresarial': [20.2, 27.5, 43.5, 36.1, 37.66, 32.3]
    },
    'ale': {
        'bajoGestion': [91.2, 67.6, 66.8, 70.0, 72.63, 78.8],
        'empresarial': [40.9, 56.5, 91.8, 74.7, 77.45, 64.7]
    }
}

# ============================================
# API ENDPOINTS PROTEGIDOS
# ============================================
@app.route('/api/user')
@login_required
def api_user():
    """Devuelve información del usuario actual"""
    return {
        'username': current_user.username,
        'email': current_user.email,
        'is_admin': current_user.is_admin
    }

@app.route('/api/data')
@login_required
def api_data():
    """Devuelve todos los datos del dashboard - REQUIERE LOGIN"""
    return {
        'periodData': PERIOD_DATA,
        'historicalData': HISTORICAL_DATA
    }

@app.route('/api/export')
@login_required
def export_data():
    """Exporta los datos con opciones personalizadas - REQUIERE LOGIN"""
    # Obtener parámetros
    formato = request.args.get('format', 'excel')
    periodos = request.args.getlist('periodos') or ['1s2025']
    perfiles = request.args.getlist('perfiles') or ['total']
    secciones = request.args.getlist('secciones') or ['resumen']

    try:
        # Preparar datos según selección
        export_data = prepare_export_data(periodos, perfiles, secciones)

        if formato == 'excel':
            return export_to_excel(export_data, periodos, perfiles)
        elif formato == 'csv':
            return export_to_csv(export_data, periodos, perfiles)
        elif formato == 'pdf':
            return export_to_pdf(export_data, periodos, perfiles)
        else:
            return {'error': 'Formato no soportado'}, 400

    except Exception as e:
        return {'error': str(e)}, 500

def prepare_export_data(periodos, perfiles, secciones):
    """Prepara los datos para exportar según las opciones seleccionadas"""
    data = {'periodos': {}, 'historico': None}

    for periodo in periodos:
        if periodo in PERIOD_DATA:
            periodo_data = PERIOD_DATA[periodo].copy()
            # Filtrar por perfiles
            filtered = {'title': periodo_data['title'], 'date': periodo_data['date']}
            for perfil in perfiles:
                if perfil in periodo_data:
                    filtered[perfil] = periodo_data[perfil]
                elif perfil == 'total' and 'total' in periodo_data:
                    filtered['total'] = periodo_data['total']
            data['periodos'][periodo] = filtered

    # Incluir datos históricos si se selecciona patrimonio
    if 'patrimonio' in secciones:
        data['historico'] = HISTORICAL_DATA

    return data

def export_to_excel(data, periodos, perfiles):
    """Exporta a Excel"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1a2744", end_color="1a2744", fill_type="solid")
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'),
                         top=Side(style='thin'), bottom=Side(style='thin'))

    first_sheet = True
    for periodo_key, periodo_data in data['periodos'].items():
        if first_sheet:
            ws = wb.active
            ws.title = periodo_data.get('title', periodo_key)[:31]
            first_sheet = False
        else:
            ws = wb.create_sheet(periodo_data.get('title', periodo_key)[:31])

        # Título
        ws['A1'] = f"BLANES CAPITAL - {periodo_data.get('title', '')}"
        ws['A1'].font = Font(bold=True, size=14)
        ws.merge_cells('A1:D1')
        ws['A2'] = f"Fecha: {periodo_data.get('date', '')}"

        # Encabezados dinámicos según perfiles seleccionados
        headers = ['Concepto']
        for perfil in perfiles:
            if perfil == 'pablo':
                headers.append('Pablo')
            elif perfil == 'ale':
                headers.append('Alejandro')
            elif perfil == 'total':
                headers.append('Total')

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin_border

        # Datos
        conceptos = [
            ('Patrimonio Total (M€)', 'patrimonioTotal'),
            ('Bajo Gestión (M€)', 'bajoGestion'),
            ('Empresarial (M€)', 'empresarial'),
            ('Rentabilidad (%)', 'rentabilidad'),
            ('Rentas Inmobiliarias (K€)', 'rentasInmob'),
            ('Cartera Financiera (M€)', 'carteraFinanciera'),
            ('Cartera Inmobiliaria (M€)', 'carteraInmobiliaria'),
            ('Inversiones Alternativas (M€)', 'alternativas'),
        ]

        row = 5
        for concepto_nombre, concepto_key in conceptos:
            ws.cell(row=row, column=1, value=concepto_nombre).border = thin_border
            col = 2
            for perfil in perfiles:
                if perfil in periodo_data and concepto_key in periodo_data[perfil]:
                    value = periodo_data[perfil][concepto_key]
                else:
                    value = '-'
                cell = ws.cell(row=row, column=col, value=value)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="right")
                col += 1
            row += 1

        ws.column_dimensions['A'].width = 30
        for c in ['B', 'C', 'D', 'E']:
            ws.column_dimensions[c].width = 15

    # Hoja de datos históricos
    if data.get('historico'):
        ws_hist = wb.create_sheet("Evolución Histórica")
        ws_hist['A1'] = "BLANES CAPITAL - Evolución Patrimonial"
        ws_hist['A1'].font = Font(bold=True, size=14)

        hist_headers = ['Perfil / Concepto'] + data['historico']['years']
        for col, header in enumerate(hist_headers, 1):
            cell = ws_hist.cell(row=3, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin_border

        row = 4
        for perfil in ['pablo', 'ale']:
            if perfil in data['historico']:
                for metric in ['bajoGestion', 'empresarial']:
                    nombre = f"{'Pablo' if perfil == 'pablo' else 'Alejandro'} - {'Bajo Gestión' if metric == 'bajoGestion' else 'Empresarial'} (M€)"
                    ws_hist.cell(row=row, column=1, value=nombre).border = thin_border
                    for col, val in enumerate(data['historico'][perfil][metric], 2):
                        cell = ws_hist.cell(row=row, column=col, value=val)
                        cell.border = thin_border
                    row += 1

        ws_hist.column_dimensions['A'].width = 30

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"BlanesCapital_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name=filename)

def export_to_csv(data, periodos, perfiles):
    """Exporta a CSV"""
    import csv

    output = BytesIO()
    writer = csv.writer(output.getvalue().decode('utf-8').encode('utf-8-sig').split())

    lines = []
    lines.append("BLANES CAPITAL - Exportación de Datos")
    lines.append(f"Fecha de exportación: {datetime.now().strftime('%d/%m/%Y %H:%M')}")
    lines.append("")

    for periodo_key, periodo_data in data['periodos'].items():
        lines.append(f"=== {periodo_data.get('title', periodo_key)} ===")
        lines.append(f"Fecha datos: {periodo_data.get('date', '')}")
        lines.append("")

        # Headers
        header_line = "Concepto"
        for perfil in perfiles:
            if perfil == 'pablo':
                header_line += ",Pablo"
            elif perfil == 'ale':
                header_line += ",Alejandro"
            elif perfil == 'total':
                header_line += ",Total"
        lines.append(header_line)

        conceptos = [
            ('Patrimonio Total (M€)', 'patrimonioTotal'),
            ('Bajo Gestión (M€)', 'bajoGestion'),
            ('Empresarial (M€)', 'empresarial'),
            ('Rentabilidad (%)', 'rentabilidad'),
            ('Rentas Inmobiliarias (K€)', 'rentasInmob'),
            ('Cartera Financiera (M€)', 'carteraFinanciera'),
            ('Cartera Inmobiliaria (M€)', 'carteraInmobiliaria'),
            ('Inversiones Alternativas (M€)', 'alternativas'),
        ]

        for concepto_nombre, concepto_key in conceptos:
            line = concepto_nombre
            for perfil in perfiles:
                if perfil in periodo_data and concepto_key in periodo_data[perfil]:
                    line += f",{periodo_data[perfil][concepto_key]}"
                else:
                    line += ",-"
            lines.append(line)
        lines.append("")

    csv_content = "\n".join(lines)
    output = BytesIO(csv_content.encode('utf-8-sig'))

    filename = f"BlanesCapital_{datetime.now().strftime('%Y%m%d_%H%M')}.csv"
    return send_file(output, mimetype='text/csv', as_attachment=True, download_name=filename)

def export_to_pdf(data, periodos, perfiles):
    """Exporta a PDF"""
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

        output = BytesIO()
        doc = SimpleDocTemplate(output, pagesize=A4, topMargin=2*cm, bottomMargin=2*cm)
        elements = []
        styles = getSampleStyleSheet()

        # Título
        title_style = ParagraphStyle('Title', parent=styles['Heading1'], fontSize=18,
                                     textColor=colors.HexColor('#1a2744'), spaceAfter=20)
        elements.append(Paragraph("BLANES CAPITAL", title_style))
        elements.append(Paragraph(f"Exportación de Datos - {datetime.now().strftime('%d/%m/%Y')}", styles['Normal']))
        elements.append(Spacer(1, 20))

        for periodo_key, periodo_data in data['periodos'].items():
            # Subtítulo del periodo
            elements.append(Paragraph(periodo_data.get('title', periodo_key), styles['Heading2']))
            elements.append(Paragraph(f"Fecha: {periodo_data.get('date', '')}", styles['Normal']))
            elements.append(Spacer(1, 10))

            # Tabla de datos
            headers = ['Concepto']
            for perfil in perfiles:
                if perfil == 'pablo':
                    headers.append('Pablo')
                elif perfil == 'ale':
                    headers.append('Alejandro')
                elif perfil == 'total':
                    headers.append('Total')

            table_data = [headers]

            conceptos = [
                ('Patrimonio Total (M€)', 'patrimonioTotal'),
                ('Bajo Gestión (M€)', 'bajoGestion'),
                ('Rentabilidad (%)', 'rentabilidad'),
                ('Rentas Inmobiliarias (K€)', 'rentasInmob'),
                ('Cartera Financiera (M€)', 'carteraFinanciera'),
                ('Cartera Inmobiliaria (M€)', 'carteraInmobiliaria'),
            ]

            for concepto_nombre, concepto_key in conceptos:
                row = [concepto_nombre]
                for perfil in perfiles:
                    if perfil in periodo_data and concepto_key in periodo_data[perfil]:
                        row.append(str(periodo_data[perfil][concepto_key]))
                    else:
                        row.append('-')
                table_data.append(row)

            table = Table(table_data, colWidths=[7*cm] + [3*cm] * (len(headers) - 1))
            table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1a2744')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('ALIGN', (0, 1), (0, -1), 'LEFT'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 10),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('GRID', (0, 0), (-1, -1), 1, colors.black),
                ('FONTSIZE', (0, 1), (-1, -1), 9),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f5f6fa')]),
            ]))
            elements.append(table)
            elements.append(Spacer(1, 30))

        doc.build(elements)
        output.seek(0)

        filename = f"BlanesCapital_{datetime.now().strftime('%Y%m%d_%H%M')}.pdf"
        return send_file(output, mimetype='application/pdf', as_attachment=True, download_name=filename)

    except ImportError:
        return {'error': 'Librería reportlab no instalada'}, 500

# ============================================
# INICIALIZACIÓN
# ============================================
def create_tables():
    """Crea las tablas y el usuario por defecto"""
    with app.app_context():
        db.create_all()

        # Crear usuario admin si no existe
        if not User.query.filter_by(username='blanes').first():
            admin = User(username='blanes', email='admin@blanescapital.com', is_admin=True)
            admin.set_password('blanes2025')
            db.session.add(admin)
            db.session.commit()
            print('Usuario admin creado: blanes / blanes2025')

# Crear tablas al iniciar
create_tables()

if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=8080)
